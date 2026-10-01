"""文件解析器：按扩展名解析常见格式为文本，供 LLM 分析。

支持：txt / md / json / csv / xlsx / docx / py / yaml / log
- txt/md/py/log 直接按文本读
- json 解析并格式化
- csv 转列表
- xlsx 用 openpyxl（可选）
- docx 用 python-docx（可选）
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from superagent.core.errors import SuperAgentError


def parse_bytes(filename: str, data: bytes) -> dict:
    """根据文件名扩展名解析字节内容，返回 {filename, kind, text, meta}。"""
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext in ("txt", "md", "py", "yaml", "yml", "log", ""):
        return _text(filename, data)
    if ext == "json":
        return _json(filename, data)
    if ext == "csv":
        return _csv(filename, data)
    if ext == "xlsx":
        return _xlsx(filename, data)
    if ext == "docx":
        return _docx(filename, data)
    # 未知类型：按文本尝试
    return _text(filename, data)


def _text(filename: str, data: bytes) -> dict:
    for enc in ("utf-8", "gbk", "utf-16"):
        try:
            return {"filename": filename, "kind": "text", "text": data.decode(enc), "meta": {"encoding": enc}}
        except UnicodeDecodeError:
            continue
    return {"filename": filename, "kind": "text", "text": data.decode("utf-8", "replace"), "meta": {"encoding": "utf-8(replace)"}}


def _json(filename: str, data: bytes) -> dict:
    obj = json.loads(data.decode("utf-8", "replace"))
    text = json.dumps(obj, ensure_ascii=False, indent=2)
    return {"filename": filename, "kind": "json", "text": text, "meta": {"json_type": type(obj).__name__}}


def _csv(filename: str, data: bytes) -> dict:
    reader = csv.DictReader(io.StringIO(data.decode("utf-8", "replace")))
    rows = list(reader)
    text = json.dumps(rows, ensure_ascii=False, indent=2)
    return {"filename": filename, "kind": "csv", "text": text, "meta": {"rows": len(rows), "columns": list(reader.fieldnames or [])}}


def _xlsx(filename: str, data: bytes) -> dict:
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError as exc:
        raise SuperAgentError("未安装 openpyxl，无法解析 xlsx") from exc
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = [[str(c) if c is not None else "" for c in row] for row in ws.iter_rows(values_only=True)]
        sheets.append({"sheet": ws.title, "rows": rows[:500]})
    text = json.dumps(sheets, ensure_ascii=False, indent=2)
    return {"filename": filename, "kind": "xlsx", "text": text, "meta": {"sheets": [s["sheet"] for s in sheets]}}


def _docx(filename: str, data: bytes) -> dict:
    try:
        import docx  # noqa: PLC0415
    except ImportError as exc:
        raise SuperAgentError("未安装 python-docx，无法解析 docx") from exc
    doc = docx.Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paras)
    return {"filename": filename, "kind": "docx", "text": text, "meta": {"paragraphs": len(paras)}}
